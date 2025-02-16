import openpyxl
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework.parsers import MultiPartParser, FormParser
from django.http import HttpResponse
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from django.http import HttpResponse
import openpyxl

from employee.permissions import IsAuthorOrReadOnly
from .models import Product, Group
from rest_framework import viewsets, generics, permissions, status
from .models import *
from.serializers import *

class OutletAPIListCreate(generics.ListCreateAPIView):
    queryset = Outlet.objects.all()
    serializer_class = OutletSerializer
    permission_classes = (IsAuthenticated,)


class OutletAPIUpdate(generics.RetrieveUpdateAPIView):
    queryset = Outlet.objects.all()
    serializer_class = OutletSerializer
    permission_classes = (IsAuthenticated,)


class OutletAPIDestroy(generics.RetrieveDestroyAPIView):
    queryset = Outlet.objects.all()
    serializer_class = OutletSerializer
    permission_classes = (IsAuthenticated,)


class ProductAPIListCreate(generics.ListCreateAPIView):
    queryset = Product.objects.all()
    serializer_class = ProductSerializer
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['group']

 
class ProductAPIUpdate(generics.RetrieveUpdateAPIView):
    queryset = Product.objects.all()
    serializer_class = ProductSerializer
    permission_classes = (IsAuthenticated,)


class ProductAPIDestroy(generics.RetrieveDestroyAPIView):
    queryset = Product.objects.all()
    serializer_class = ProductSerializer
    permission_classes = (IsAuthenticated,)


class OrderAPIListCreate(generics.ListCreateAPIView):
    queryset = Order.objects.all()
    serializer_class = OrderSerializer
    permission_classes = (IsAuthenticated,)


class OrderAPIUpdate(generics.RetrieveUpdateAPIView):
    queryset = Order.objects.all()
    serializer_class = OrderSerializer
    permission_classes = (IsAuthenticated,)


class ClientAPIListCreate(generics.ListCreateAPIView):
    queryset = Client.objects.all()
    serializer_class = ClientSerializer
    permission_classes = (IsAuthenticated,)


class ClientAPIUpdate(generics.RetrieveUpdateAPIView):
    queryset = Client.objects.all()
    serializer_class = ClientSerializer
    permission_classes = (IsAuthenticated,)


class ClientAPIDestroy(generics.RetrieveDestroyAPIView):
    queryset = Client.objects.all()
    serializer_class = ClientSerializer
    permission_classes = (IsAuthenticated,)


class GroupAPIListCreate(generics.ListCreateAPIView):
    queryset = Group.objects.all()
    serializer_class = GroupSerializer


class GroupAPIUpdateDestroy(generics.RetrieveUpdateDestroyAPIView):
    queryset = Group.objects.all()
    serializer_class = GroupSerializer



class ProductImportExportView(APIView):
    parser_classes = (MultiPartParser, FormParser)

    def post(self, request):
        """Импорт и обновление продуктов из Excel файла"""
        file = request.FILES.get('file')
        if not file:
            return Response({"error": "Файл не предоставлен"}, status=status.HTTP_400_BAD_REQUEST)
        try:
            workbook = openpyxl.load_workbook(file)
            sheet = workbook.active
            for row in sheet.iter_rows(min_row=2, values_only=True):
                group_name, name, price, count, article, description = row
                group, _ = Group.objects.get_or_create(name=group_name)
                product, created = Product.objects.update_or_create(
                    article=article,
                    defaults={
                        "group": group,
                        "name": name,
                        "price": price,
                        "count": count,
                        "description": description,
                    }
                )

            return Response({"message": "Продукты успешно импортированы/обновлены"}, status=status.HTTP_200_OK)

        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

    def get(self, request):
        """Экспорт продуктов в стилизованный Excel-файл"""
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Продукты"

        # Определяем стили
        header_font = Font(bold=True, color="FFFFFF")  # Белый текст
        header_fill = PatternFill(start_color="007bff", end_color="007bff", fill_type="solid")  # Синий фон
        center_alignment = Alignment(horizontal="center", vertical="center")  # Выравнивание по центру
        border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"),
                        bottom=Side(style="thin"))

        # Заголовки
        headers = ["Группа", "Название", "Цена", "Количество", "Артикул", "Описание"]
        sheet.append(headers)

        # Применяем стили к заголовкам
        for col_num, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
            cell.border = border

        # Данные из базы
        products = Product.objects.select_related('group').all()
        for row_num, product in enumerate(products, start=2):
            sheet.append([
                product.group.name,
                product.name,
                product.price,
                product.count,
                product.article,
                product.description,
            ])

            # Чередование цветов строк (серый фон для четных строк)
            if row_num % 2 == 0:
                fill = PatternFill(start_color="f2f2f2", end_color="f2f2f2", fill_type="solid")
                for col_num in range(1, len(headers) + 1):
                    sheet.cell(row=row_num, column=col_num).fill = fill

            # Форматируем числовые значения
            sheet.cell(row=row_num, column=3).number_format = '#,##0.00'  # Цена
            sheet.cell(row=row_num, column=4).number_format = '0'  # Количество

        # Автоширина колонок
        for col in sheet.columns:
            max_length = 0
            col_letter = col[0].column_letter  # Получаем букву колонки
            for cell in col:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            sheet.column_dimensions[col_letter].width = max_length + 2

        # Создаем HTTP-ответ с файлом Excel
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=products.xlsx'

        # Сохраняем книгу в ответ
        workbook.save(response)
        return response
